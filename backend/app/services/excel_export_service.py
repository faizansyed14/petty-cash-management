"""Excel export matching ALPHA DATA RECRUITMENT voucher register sheet."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.db.models.voucher import SettlementType
from app.services.voucher_service import list_vouchers


HEADERS = [
    "S.NO:",
    "VOUCHER NO:",
    "NAME",
    "AMOUNT",
    "NOTES",
    "DEPOSITED",
    "CASH COLLECTION",
    "DATE",
    "FILE NUMBER",
]


def _note_text(v: Voucher) -> str:
    return v.details or ""


def build_voucher_register_excel(
    db: Session,
    date_from: date | None = None,
    date_to: date | None = None,
    status: str | None = None,
    search: str | None = None,
) -> bytes:
    from app.db.models.voucher import VoucherStatus

    single = None
    statuses = None
    if status:
        parts = [p.strip() for p in status.split(",") if p.strip()]
        parsed = [VoucherStatus(p) for p in parts]
        if len(parsed) == 1:
            single = parsed[0]
        else:
            statuses = parsed

    vouchers, _ = list_vouchers(
        db,
        status=single,
        statuses=statuses,
        search=search,
        date_from=date_from,
        date_to=date_to,
        skip=0,
        limit=5000,
    )
    vouchers = sorted(vouchers, key=lambda v: (v.settlement_date or date.min, v.voucher_date))

    wb = Workbook()
    ws = wb.active
    ws.title = "Vouchers"

    # Title row — ALPHA DATA RECRUITMENT
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = "ALPHA DATA RECRUITMENT"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="008000")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    dashed = Border(
        left=Side(style="dashed", color="000000"),
        right=Side(style="dashed", color="000000"),
        top=Side(style="dashed", color="000000"),
        bottom=Side(style="dashed", color="000000"),
    )
    header_font = Font(name="Calibri", size=11, bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = header_font
        cell.alignment = header_align
        cell.border = dashed
    ws.row_dimensions[2].height = 22

    body_font = Font(name="Calibri", size=11)
    for idx, v in enumerate(vouchers, start=1):
        row = idx + 2
        deposited = ""
        cash_collection = ""
        if v.settlement_type == SettlementType.DEPOSITED:
            deposited = v.bank_name or "Deposited"
        elif v.settlement_type == SettlementType.CASH_COLLECTION:
            cash_collection = "Cash Collection"

        values = [
            idx,
            v.voucher_no,
            v.employee.name if v.employee else "",
            float(v.total_amount),
            _note_text(v),
            deposited,
            cash_collection,
            v.settlement_date,
            v.file_number or "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value if value is not None else "")
            cell.font = body_font
            cell.alignment = Alignment(
                horizontal="center" if col in (1, 4, 8, 9) else "left",
                vertical="center",
            )
            if col == 4:
                cell.number_format = "#,##0.00"
            if col == 8 and v.settlement_date:
                cell.number_format = "DD/MM/YYYY"

    widths = [8, 16, 28, 12, 28, 28, 18, 14, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
